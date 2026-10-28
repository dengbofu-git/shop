"""Bounded spreadsheet parsing. No Django, ORM, network or model calls."""
import csv
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
import json
import re
import zipfile

from .catalog_constants import TEA_CATEGORIES

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 500
MAX_COLUMNS = 20
PARSER_VERSION = 'products-sheet-v2'
HEADERS = {'商品名称': 'name', '分类': 'catalog', '规格名称': 'sku_name', '品牌': 'brand',
           '商品介绍': 'caption', '售价': 'price', '库存': 'stock',
           '市场价': 'market_price', '币种': 'currency'}
ALIASES = dict(HEADERS, **{name: name for name in HEADERS.values()})
REQUIRED = {'name', 'catalog', 'price', 'stock'}
TEMPLATE_HEADERS = list(HEADERS)
TEMPLATE_ROW = ['示例红茶', '红茶', '150g袋装', '示例品牌', '请替换为真实商品介绍',
                '38.00', '100', '38.00', 'CNY']


class ImportFileError(ValueError):
    pass


def _text(value):
    if value is None:
        return ''
    if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
        raise ImportFileError('仅支持文本和数值单元格，请将日期等内容转换为文本')
    value = str(value).strip()
    if len(value) > 1000 or any(ord(c) < 32 and c not in '\r\n\t' for c in value):
        raise ImportFileError('单元格内容过长或含有不可见控制字符')
    return value


def _csv_rows(content):
    for encoding in ('utf-8-sig', 'gb18030'):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportFileError('CSV 编码不支持，请另存为 UTF-8 CSV')
    try:
        for line, row in enumerate(csv.reader(StringIO(text, newline=''), strict=True), 1):
            if line > MAX_ROWS + 1:
                raise ImportFileError('每批最多 500 行，请拆分文件')
            yield line, row
    except csv.Error:
        raise ImportFileError('CSV 格式错误，请检查引号、分隔符并使用逗号分隔') from None


def _xlsx_rows(content):
    try:
        from defusedxml import ElementTree
        from openpyxl import load_workbook
        from openpyxl.xml import DEFUSEDXML
    except ImportError:
        raise ImportFileError('XLSX 解析依赖未安装，请联系管理员安装项目依赖') from None
    if not DEFUSEDXML:
        raise ImportFileError('XLSX 安全解析器未启用')
    workbook = None
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if (len(entries) > 1000 or len({e.filename for e in entries}) != len(entries)
                    or sum(e.file_size for e in entries) > 30 * 1024 * 1024
                    or any(e.file_size > 10 * 1024 * 1024 or e.flag_bits & 1 for e in entries)):
                raise ImportFileError('XLSX 解压内容超限或不支持加密文件')
            for entry in entries:
                name = entry.filename.lower()
                if 'vbaproject' in name or name.startswith(('xl/externallinks/', 'xl/embeddings/')):
                    raise ImportFileError('不支持宏、外部链接或嵌入对象，请使用纯数据模板')
                if name.endswith(('.xml', '.rels')):
                    # Validate every XML, not just the worksheet consumed by openpyxl.
                    with archive.open(entry) as stream:
                        for _, element in ElementTree.iterparse(stream, events=('end',)):
                            local = element.tag.rsplit('}', 1)[-1]
                            if local == 'mergeCell':
                                raise ImportFileError('不支持合并单元格，请按模板一行填写一个规格')
                            if local == 'row' and int(element.attrib.get('r', '0')) > MAX_ROWS + 1:
                                raise ImportFileError('每批最多 500 行，请拆分文件')
                            if local == 'c':
                                address = element.attrib.get('r', '')
                                match = re.fullmatch(r'([A-Z]+)([1-9]\d*)', address)
                                if match:
                                    column = 0
                                    for letter in match[1]:
                                        column = column * 26 + ord(letter) - 64
                                    if column > MAX_COLUMNS or int(match[2]) > MAX_ROWS + 1:
                                        raise ImportFileError('表格超出 500 行或 20 列，请移除多余格式并拆分文件')
                            if local == 'Relationship' and element.attrib.get('TargetMode') == 'External':
                                raise ImportFileError('不支持外部链接，请移除超链接和外部引用')
                            element.clear()
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if len(workbook.sheetnames) != 1 or not workbook.worksheets:
            raise ImportFileError('首期只支持一个工作表，请拆分后上传')
        sheet = workbook.worksheets[0]
        sheet.reset_dimensions()  # Do not trust forged worksheet dimensions.
        for line, cells in enumerate(sheet.iter_rows(), 1):
            if line > MAX_ROWS + 1:
                raise ImportFileError('每批最多 500 行，请拆分文件')
            if any(cell.data_type in ('f', 'e') for cell in cells):
                raise ImportFileError('第 {} 行含公式或错误单元格，请粘贴为数值后上传'.format(line))
            yield line, [cell.value for cell in cells]
    except ImportFileError:
        raise
    except Exception:
        raise ImportFileError('XLSX 损坏或格式不支持，请使用下载的模板另存后上传') from None
    finally:
        if workbook is not None:
            workbook.close()


def _money(value, label, errors):
    # No exponent notation, separators, currency symbols or silent rounding.
    if not re.fullmatch(r'\d{1,8}(?:\.\d{1,2})?', value):
        errors.append(label + '须为 0～99999999.99 的金额，最多两位小数')
        return value
    try:
        return format(Decimal(value).quantize(Decimal('.01')), 'f')
    except InvalidOperation:
        errors.append(label + '格式不正确')
        return value


def normalize_row(raw):
    """Return a typed candidate plus errors; never turn malformed stock into zero."""
    if 'cost_price' in raw or '进价' in raw:
        raise ImportFileError('商品导入不采集内部成本，请删除进价列并使用新模板')
    data = {name: _text(raw.get(name)) for name in HEADERS.values()}
    errors = []
    for key, limit, label in [('name', 50, '商品名称'), ('sku_name', 50, '规格名称'),
                              ('brand', 20, '品牌'), ('caption', 100, '商品介绍')]:
        if len(data[key]) > limit:
            errors.append('{}最多 {} 字'.format(label, limit))
    if not data['name']:
        errors.append('商品名称必填')
    data['sku_name'] = data['sku_name'] or data['name']
    if data['catalog'] not in TEA_CATEGORIES:
        errors.append('分类须为：' + '、'.join(TEA_CATEGORIES))
    data['currency'] = data['currency'].upper() or 'CNY'
    if data['currency'] != 'CNY':
        errors.append('首期仅支持人民币 CNY，请先确认币种和价格')
    data['price'] = _money(data['price'], '售价', errors)
    data['market_price'] = _money(data['market_price'] or data['price'], '市场价', errors)
    if not re.fullmatch(r'\d{1,10}', data['stock']) or int(data['stock']) > 2147483647:
        errors.append('库存须为 0～2147483647 的整数')
    else:
        data['stock'] = int(data['stock'])
    return data, errors


def parse_file(filename, content):
    if not isinstance(content, bytes) or not 0 < len(content) <= MAX_BYTES:
        raise ImportFileError('文件不能为空，且不得超过 5 MB')
    extension = filename.rsplit('.', 1)[-1].lower()
    if extension not in ('csv', 'xlsx'):
        raise ImportFileError('仅支持 .xlsx 和 .csv 文件')
    source = iter(_csv_rows(content) if extension == 'csv' else _xlsx_rows(content))
    first = next(source, None)
    if first is None:
        raise ImportFileError('文件没有表头和商品数据')
    columns = [_text(value) for value in first[1]]
    # Ignore only trailing empty columns commonly saved by Excel.
    while columns and not columns[-1]:
        columns.pop()
    if any(column.lower() in ('进价', 'cost_price') for column in columns):
        raise ImportFileError('商品导入不采集内部成本，请删除进价列并使用新模板')
    if len(columns) > MAX_COLUMNS or not columns:
        raise ImportFileError('表头列数不正确，请使用模板')
    headers = [ALIASES.get(value.lower(), ALIASES.get(value)) for value in columns]
    if None in headers or len(set(headers)) != len(headers) or not REQUIRED.issubset(headers):
        raise ImportFileError('表头不正确：商品名称、分类、售价、库存必填；请使用模板且不要重复列名')
    rows = []
    groups, seen = {}, {}
    for line, values in source:
        if len(values) > MAX_COLUMNS:
            raise ImportFileError('第 {} 行列数超限'.format(line))
        values = [_text(value) for value in values]
        if any(value.startswith('=') for value in values):
            raise ImportFileError('第 {} 行含公式，请粘贴为数值后上传'.format(line))
        if not any(values):
            continue
        if any(values[len(headers):]):
            raise ImportFileError('第 {} 行有未命名列的数据，请检查表头'.format(line))
        data, errors = normalize_row(dict(zip(headers, values)))
        key = (data['name'].casefold(), data['sku_name'].casefold())
        if key in seen:
            errors.append('与第 {} 行商品和规格重复'.format(seen[key]))
        seen[key] = line
        group = (data['catalog'], data['brand'])
        name = data['name'].casefold()
        if name in groups and groups[name] != group:
            errors.append('同名商品的分类和品牌必须一致')
        groups[name] = group
        rows.append({'row_number': line, 'data': data, 'errors': errors})
    if not rows:
        raise ImportFileError('文件没有商品数据，请从第 2 行填写')
    return rows


def canonical(rows):
    return json.dumps(rows, sort_keys=True, ensure_ascii=False, separators=(',', ':'))
