"""Import contract tested in disposable application databases, never business data."""
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from io import BytesIO
import json
from threading import Barrier
import uuid
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import close_old_connections, connection
from django.test import Client, SimpleTestCase, TestCase, TransactionTestCase
from django.utils import timezone
from openpyxl import Workbook

from tea_saas.utils import jwt_encode
from users.models import UserProfile
from inventory.models import Inventory, InventoryLog, Warehouse
from inventory.services import InventoryService
from .models import Brand, Catalog, ProductImportJob, SKU, SPU
from .product_import_parser import ImportFileError, MAX_BYTES, parse_file


CSV = '商品名称,分类,规格名称,品牌,商品介绍,售价,库存,市场价,币种\n山间红茶,红茶,150g袋装,茶园,花果香,38.50,12,45,CNY\n山间红茶,红茶,500g礼盒,茶园,礼盒装,98.00,5,128,CNY\n'


def xlsx(rows, mutate=None):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    if mutate:
        mutate(book)
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


class ProductImportParserTests(SimpleTestCase):
    def test_csv_encodings_and_typed_money(self):
        for encoding in ('utf-8', 'utf-8-sig', 'gb18030'):
            with self.subTest(encoding=encoding):
                rows = parse_file('商品.CSV', CSV.encode(encoding))
                self.assertEqual(len(rows), 2)
                self.assertEqual(rows[0]['errors'], [])
                self.assertEqual(rows[0]['data']['price'], '38.50')
                self.assertEqual(rows[0]['data']['stock'], 12)
                self.assertEqual(rows[1]['row_number'], 3)

    def test_xlsx_and_english_headers(self):
        rows = parse_file('a.xlsx', xlsx([['name', 'catalog', 'price', 'stock'], ['红茶', '红茶', 12.5, 0]]))
        self.assertEqual(rows[0]['errors'], [])
        self.assertEqual(rows[0]['data']['price'], '12.50')
        self.assertEqual(rows[0]['data']['sku_name'], '红茶')

    def test_row_errors_do_not_coerce_values(self):
        text = '商品名称,分类,售价,库存,币种\n测试茶,错误茶类,1.234,错误,USD\n'
        row = parse_file('a.csv', text.encode())[0]
        self.assertGreaterEqual(len(row['errors']), 4)
        self.assertEqual(row['data']['stock'], '错误')
        self.assertEqual(row['data']['price'], '1.234')

    def test_numeric_ranges_required_fields_and_lengths(self):
        cases = [('售价', '-1'), ('售价', 'NaN'), ('售价', 'Infinity'), ('售价', '1e2'),
                 ('售价', '100000000'), ('库存', '-1'), ('库存', '1.5'), ('库存', '2147483648'),
                 ('商品名称', ''), ('商品名称', '茶' * 51), ('品牌', '茶' * 21), ('商品介绍', '茶' * 101)]
        for field, value in cases:
            row = {'商品名称': '测试', '分类': '红茶', '售价': '1', '库存': '0', '品牌': '', '商品介绍': ''}
            row[field] = value
            with self.subTest(field=field, value=value):
                parsed = parse_file('a.xlsx', xlsx([list(row), list(row.values())]))
                self.assertTrue(parsed[0]['errors'])

    def test_duplicate_variant_and_conflicting_spu_metadata(self):
        rows = parse_file('a.csv', (CSV + '山间红茶,绿茶,150g袋装,别家,,1,2,,\n').encode())
        self.assertTrue(any('重复' in error for error in rows[-1]['errors']))
        self.assertTrue(any('一致' in error for error in rows[-1]['errors']))

    def test_unknown_duplicate_missing_headers_and_unheaded_data(self):
        for text in ['商品名称,售价,库存\n茶,1,0\n', '商品名称,分类,售价,库存,库存\n茶,红茶,1,0,0\n',
                     '商品名称,分类,售价,库存,商家ID\n茶,红茶,1,0,1\n',
                     '商品名称,分类,售价,库存\n茶,红茶,1,0,多余\n']:
            with self.subTest(text=text), self.assertRaises(ImportFileError):
                parse_file('a.csv', text.encode())

    def test_file_format_size_row_and_column_limits(self):
        for name, content in [('a.pdf', b'%PDF-'), ('a.xlsx', b'broken'), ('a.csv', b''),
                              ('a.csv', b'x' * (MAX_BYTES + 1)),
                              ('a.csv', ('name,catalog,price,stock\n' + 'Tea,红茶,1,0\n' * 501).encode()),
                              ('a.xlsx', xlsx([['name', 'catalog', 'price', 'stock'], ['Tea', '红茶', 1, 0]],
                                  lambda book: setattr(book.active.cell(2, 21), 'value', 'extra')))]:
            with self.subTest(name=name, size=len(content)), self.assertRaises(ImportFileError):
                parse_file(name, content)

    def test_formula_multiple_sheets_merged_and_dates_are_rejected(self):
        from datetime import datetime
        base = [['name', 'catalog', 'price', 'stock'], ['Tea', '红茶', 1, 0]]
        for mutate in [lambda book: setattr(book.active['C2'], 'value', '=1+1'),
                       lambda book: book.create_sheet('extra'),
                       lambda book: book.active.merge_cells('A2:B2'),
                       lambda book: setattr(book.active['C2'], 'value', datetime(2026, 1, 1))]:
            with self.assertRaises(ImportFileError):
                parse_file('a.xlsx', xlsx(base, mutate))
        with self.assertRaises(ImportFileError):
            parse_file('a.csv', b'name,catalog,price,stock\n=SUM(1),black,1,0')

    def test_zip_xml_entity_is_rejected(self):
        import zipfile
        stream = BytesIO()
        with zipfile.ZipFile(stream, 'w') as archive:
            archive.writestr('bad.xml', '<!DOCTYPE x [<!ENTITY a "unsafe">]><x>&a;</x>')
        with self.assertRaises(ImportFileError):
            parse_file('a.xlsx', stream.getvalue())

    def test_internal_cost_columns_are_explicitly_rejected(self):
        for header in ('进价', 'cost_price'):
            with self.subTest(header=header), self.assertRaisesMessage(ImportFileError, '不采集内部成本'):
                parse_file('a.csv', ('name,catalog,price,stock,' + header + '\n茶品,红茶,1,0,10\n').encode())


class ImportHttpMixin:
    def auth(self, user=None):
        token = jwt_encode({'id': (user or self.merchant).pk})
        if isinstance(token, bytes):
            token = token.decode()
        return {'HTTP_AUTHORIZATION': 'Bearer ' + token}

    def upload(self, content=None, key=None, user=None, filename='商品.csv', client=None):
        return (client or self.client).post('/v1/goods/merchant/imports', {
            'file': SimpleUploadedFile(filename, content if content is not None else CSV.encode()),
            'request_key': str(key or uuid.uuid4())}, **self.auth(user))

    def confirm(self, job, user=None, client=None, **extra):
        return (client or self.client).post('/v1/goods/merchant/imports/{}/commit'.format(job['id']),
            json.dumps(dict(review_digest=job['review_digest'], **extra)), content_type='application/json', **self.auth(user))


class ProductImportHttpTests(ImportHttpMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.merchant = UserProfile.objects.create(username='import-a', user_type='merchant', is_active=True)
        cls.other = UserProfile.objects.create(username='import-b', user_type='merchant', is_active=True)
        cls.buyer = UserProfile.objects.create(username='import-buy', is_active=True)
        cls.inactive = UserProfile.objects.create(username='import-off', user_type='merchant', is_active=False)

    def test_auth_active_role_method_and_cross_merchant_boundaries(self):
        unauthenticated = self.client.get('/v1/goods/merchant/imports/template')
        self.assertEqual(unauthenticated.status_code, 401)
        self.assertEqual(unauthenticated['Cache-Control'], 'no-store')
        self.assertEqual(self.upload(user=self.buyer).status_code, 403)
        self.assertEqual(self.upload(user=self.inactive).status_code, 401)
        self.assertEqual(self.client.get('/v1/goods/merchant/imports', **self.auth()).status_code, 405)
        job = self.upload().json()['data']
        self.assertEqual(self.confirm(job, user=self.other).status_code, 404)
        self.assertEqual(self.client.get('/v1/goods/merchant/imports/' + job['id'], **self.auth(self.other)).status_code, 404)
        self.assertEqual(SKU.objects.count(), 0)

    def test_upload_is_preview_only_then_commit_is_idempotent_and_unlaunched(self):
        response = self.upload()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Cache-Control'], 'no-store')
        job = response.json()['data']
        self.assertEqual(job['error_count'], 0)
        self.assertFalse(SPU.objects.exists())
        first = self.confirm(job)
        self.assertEqual(first.status_code, 200, first.content)
        second = self.confirm(job)
        self.assertEqual(first.json(), second.json())
        self.assertEqual(first.json()['data']['result']['spu_count'], 1)
        self.assertEqual(SPU.objects.get().merchant, self.merchant)
        self.assertEqual(SKU.objects.count(), 2)
        self.assertFalse(SKU.objects.filter(is_launched=True).exists())
        self.assertEqual(str(SKU.objects.order_by('pk').first().price), '38.50')
        self.assertEqual(list(Inventory.objects.order_by('sku_id').values_list('quantity', 'reserved_quantity', 'available_quantity')), [(12, 0, 12), (5, 0, 5)])
        self.assertEqual(InventoryLog.objects.count(), 2)
        self.assertEqual(list(SKU.objects.order_by('pk').values_list('stock', flat=True)), [12, 5])

    def test_upload_retry_reuses_job_and_key_cannot_bind_another_file(self):
        key = uuid.uuid4()
        first = self.upload(key=key).json()['data']
        self.assertEqual(self.upload(key=key).json()['data']['id'], first['id'])
        self.confirm(first)
        self.assertEqual(self.upload(key=key).json()['data']['status'], 'committed')
        self.assertEqual(self.upload(content=CSV.replace('12', '13').encode(), key=key).status_code, 409)
        self.assertEqual(ProductImportJob.objects.count(), 1)

    def test_invalid_rows_block_entire_batch(self):
        job = self.upload(content=CSV.replace('38.50', 'wrong').encode()).json()['data']
        self.assertEqual(job['status'], 'invalid')
        self.assertEqual(job['error_count'], 1)
        self.assertEqual(self.confirm(job).status_code, 409)
        self.assertFalse(SPU.objects.exists())

    def test_commit_accepts_only_confirmed_digest(self):
        job = self.upload().json()['data']
        changed = dict(job, review_digest='0' * 64)
        self.assertEqual(self.confirm(changed).status_code, 409)
        self.assertEqual(self.confirm(job, merchant_id=self.other.pk).status_code, 400)
        self.assertEqual(self.confirm(job, rows=[]).status_code, 400)
        self.assertFalse(SPU.objects.exists())

    def test_existing_products_are_not_overwritten_and_rechecked_at_commit(self):
        job = self.upload().json()['data']
        catalog = Catalog.objects.create(name='红茶')
        old = SPU.objects.create(name='山间红茶', catalog=catalog, merchant=self.merchant)
        self.assertEqual(self.confirm(job).status_code, 409)
        self.assertEqual(SPU.objects.get(pk=old.pk).name, '山间红茶')
        new = self.upload().json()['data']
        self.assertEqual(new['error_count'], 2)
        # Same names belonging to another merchant do not block a merchant's own import.
        self.assertEqual(self.upload(user=self.other).json()['data']['error_count'], 0)

    def test_expiry_and_disabled_warehouse_block_writes(self):
        job = self.upload().json()['data']
        ProductImportJob.objects.filter(pk=job['id']).update(expires_at=timezone.now() - timedelta(seconds=1))
        self.assertEqual(self.confirm(job).status_code, 409)
        another = self.upload().json()['data']
        Warehouse.objects.create(code='DEFAULT', name='默认', address='测试', is_active=False)
        self.assertEqual(self.confirm(another).status_code, 409)
        self.assertFalse(SPU.objects.exists())

    def test_failure_rolls_back_products_stock_receipt_and_can_retry(self):
        job = self.upload().json()['data']
        real_adjust = InventoryService.adjust_stock
        calls = []
        def fail_second(*args, **kwargs):
            calls.append(args)
            if len(calls) == 2:
                raise RuntimeError('sensitive fixture diagnostic must not escape')
            return real_adjust(*args, **kwargs)
        with patch.object(InventoryService, 'adjust_stock', side_effect=fail_second):
            response = self.confirm(job)
        self.assertEqual(response.status_code, 500)
        self.assertNotIn('sensitive', response.content.decode())
        for model in (SPU, SKU, Brand, Catalog, Inventory, InventoryLog, Warehouse):
            self.assertFalse(model.objects.exists(), model.__name__)
        self.assertEqual(ProductImportJob.objects.get(pk=job['id']).status, 'ready')
        self.assertEqual(self.confirm(job).status_code, 200)

    def test_tampered_persisted_candidate_and_deactivated_identity_are_rejected(self):
        job = self.upload().json()['data']
        stored = ProductImportJob.objects.get(pk=job['id'])
        stored.rows[0]['data']['stock'] = 1000
        stored.save(update_fields=['rows'])
        self.assertEqual(self.confirm(job).status_code, 409)
        UserProfile.objects.filter(pk=self.merchant.pk).update(is_active=False)
        self.assertEqual(self.confirm(job).status_code, 401)

    def test_template_roundtrips_and_xlsx_http_import(self):
        for kind in ('csv', 'xlsx'):
            response = self.client.get('/v1/goods/merchant/imports/template?format=' + kind, **self.auth())
            self.assertEqual(response.status_code, 200)
            self.assertIn('attachment;', response['Content-Disposition'])
            self.assertEqual(parse_file('template.' + kind, response.content)[0]['errors'], [])
            self.assertNotIn('cost_price', parse_file('template.' + kind, response.content)[0]['data'])
        response = self.upload(xlsx([['name', 'catalog', 'price', 'stock'], ['绿茶礼盒', '绿茶', 12.5, 4]]), filename='test.xlsx')
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(self.confirm(response.json()['data']).status_code, 200)

    def test_legacy_preview_hides_internal_cost_and_cannot_be_committed(self):
        job = self.upload().json()['data']
        self.assertNotIn('cost_price', job['rows'][0]['data'])
        stored = ProductImportJob.objects.get(pk=job['id'])
        stored.parser_version = 'products-sheet-v1'
        stored.rows[0]['data']['cost_price'] = '98765.43'
        stored.save(update_fields=['parser_version', 'rows'])
        response = self.client.get('/v1/goods/merchant/imports/' + job['id'], **self.auth())
        self.assertNotIn('cost_price', response.content.decode())
        self.assertNotIn('98765.43', response.content.decode())
        self.assertTrue(response.json()['data']['needs_reupload'])
        self.assertEqual(self.confirm(job).status_code, 409)
        self.assertFalse(SKU.objects.exists())

    def test_buyer_catalog_and_detail_do_not_return_internal_cost(self):
        job = self.upload().json()['data']
        self.assertEqual(self.confirm(job).status_code, 200)
        sku = SKU.objects.order_by('pk').first()
        sku.cost_price = '98765.43'
        sku.is_launched = True
        sku.save(update_fields=['cost_price', 'is_launched'])
        for path in ['/v1/goods/catalogs/' + str(sku.spu.catalog_id), '/v1/goods/detail/' + str(sku.pk)]:
            response = self.client.get(path)
            self.assertEqual(response.status_code, 200)
            self.assertNotIn('cost_price', response.content.decode())
            self.assertNotIn('98765.43', response.content.decode())

    def test_upload_rejects_identity_fields_oversize_and_bad_files_without_job(self):
        response = self.client.post('/v1/goods/merchant/imports', {
            'file': SimpleUploadedFile('a.csv', CSV.encode()), 'request_key': str(uuid.uuid4()),
            'merchant_id': self.other.pk}, **self.auth())
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.upload(b'x' * (MAX_BYTES + 1)).status_code, 413)
        self.assertEqual(self.upload(b'invalid xlsx', filename='bad.xlsx').status_code, 400)
        self.assertEqual(self.upload(key='invalid-uuid').status_code, 400)
        self.assertFalse(ProductImportJob.objects.exists())

    def test_maximum_batch_creates_all_500_specs(self):
        import time
        content = ('name,catalog,sku_name,price,stock\n' + ''.join(
            '大批次红茶,红茶,规格{},1.25,2\n'.format(i) for i in range(500))).encode()
        started = time.monotonic()
        job = self.upload(content).json()['data']
        self.assertEqual(job['row_count'], 500)
        response = self.confirm(job)
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['data']['result']['sku_count'], 500)
        self.assertEqual(SKU.objects.count(), 500)
        self.assertEqual(Inventory.objects.count(), 500)
        self.assertEqual(InventoryLog.objects.count(), 500)
        self.assertEqual(SPU.objects.count(), 1)
        print('IMPORT_500_SPECS_SECONDS={:.3f}'.format(time.monotonic() - started))


class ProductImportConcurrentTests(ImportHttpMixin, TransactionTestCase):
    def setUp(self):
        self.merchant = UserProfile.objects.create(username='import-race', user_type='merchant', is_active=True)

    def test_parallel_confirm_creates_one_receipt_and_one_set_of_skus(self):
        if connection.vendor != 'mysql':
            self.skipTest('Requires disposable MySQL row locks')
        job = self.upload().json()['data']
        barrier = Barrier(2)
        def confirm():
            close_old_connections()
            try:
                barrier.wait(timeout=15)
                response = self.confirm(job, client=Client())
                return response.status_code, response.json()['data']['result']
            finally:
                close_old_connections()
        with ThreadPoolExecutor(max_workers=2) as executor:
            results = list(executor.map(lambda _: confirm(), range(2)))
        self.assertEqual(results[0], results[1])
        self.assertEqual(results[0][0], 200)
        self.assertEqual(SKU.objects.count(), 2)
        self.assertEqual(SPU.objects.count(), 1)

    def test_two_ready_jobs_with_same_names_cannot_both_commit(self):
        if connection.vendor != 'mysql':
            self.skipTest('Requires disposable MySQL row locks')
        jobs = [self.upload().json()['data'], self.upload().json()['data']]
        barrier = Barrier(2)
        def confirm(job):
            close_old_connections()
            try:
                barrier.wait(timeout=15)
                return self.confirm(job, client=Client()).status_code
            finally:
                close_old_connections()
        with ThreadPoolExecutor(max_workers=2) as executor:
            statuses = list(executor.map(confirm, jobs))
        self.assertEqual(sorted(statuses), [200, 409])
        self.assertEqual(SKU.objects.count(), 2)
